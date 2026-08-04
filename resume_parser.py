import os
import pdfplumber
import docx


class ParseError(Exception):
    pass


def extract_text(file_path: str) -> str:
   
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_from_pdf(file_path)
    elif ext == ".docx":
        return _extract_from_docx(file_path)
    elif ext == ".txt":
        return _extract_from_txt(file_path)
    else:
        raise ParseError(
            f"Unsupported file type '{ext}'. Please upload a PDF, DOCX, or TXT file."
        )


def _extract_from_pdf(file_path: str) -> str:
    text_chunks = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)
    except Exception as e:
        raise ParseError(f"Could not read PDF file: {e}")

    text = "\n".join(text_chunks).strip()
    if not text:
        raise ParseError(
            "No readable text found in this PDF. It may be a scanned image "
            "please upload a text-based PDF or a DOCX file instead."
        )
    return text


def _extract_from_docx(file_path: str) -> str:
    try:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs]

        
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        paragraphs.append(cell.text)

        text = "\n".join(paragraphs).strip()
    except Exception as e:
        raise ParseError(f"Could not read DOCX file: {e}")

    if not text:
        raise ParseError("No readable text found in this DOCX file.")
    return text


def _extract_from_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
    except Exception as e:
        raise ParseError(f"Could not read TXT file: {e}")

    if not text:
        raise ParseError("This TXT file appears to be empty.")
    return text
